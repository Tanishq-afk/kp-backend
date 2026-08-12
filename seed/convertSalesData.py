#!/usr/bin/env python3
"""
One-time convert Data/saleslist.xlsx + Data/salesreturnlist.xlsx -> JSON.

Requires openpyxl (pip3 install --user openpyxl).

Usage:
    python3 seed/convertSalesData.py
    (writes seed/data/saleslist.json and seed/data/salesreturnlist.json)
"""
import datetime
import json
from pathlib import Path

import openpyxl

DATA_DIR = Path("Data")
OUT_DIR = Path("seed/data")


def combine_dt(d, t):
    if not isinstance(d, (datetime.datetime, datetime.date)):
        return None
    base_date = d.date() if isinstance(d, datetime.datetime) else d
    if isinstance(t, datetime.time):
        return datetime.datetime.combine(base_date, t).isoformat()
    if isinstance(d, datetime.datetime):
        return d.isoformat()
    return datetime.datetime.combine(base_date, datetime.time()).isoformat()


def convert_sales():
    wb = openpyxl.load_workbook(DATA_DIR / "saleslist.xlsx", read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows_iter)]
    out = []
    for row in rows_iter:
        rec = dict(zip(header, row))
        rec["billDateTime"] = combine_dt(rec.get("Billdate"), rec.get("BillTime"))
        rec.pop("Billdate", None)
        rec.pop("BillTime", None)
        out.append(rec)
    wb.close()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with (OUT_DIR / "saleslist.json").open("w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, default=str)
    print(f"Wrote {len(out)} rows -> {OUT_DIR / 'saleslist.json'}")


def convert_returns():
    wb = openpyxl.load_workbook(DATA_DIR / "salesreturnlist.xlsx", read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows_iter)]
    out = []
    for row in rows_iter:
        rec = dict(zip(header, row))
        rd = rec.get("Returndate")
        rec["returnDateTime"] = rd.isoformat() if isinstance(rd, (datetime.datetime, datetime.date)) else None
        rec.pop("Returndate", None)
        out.append(rec)
    wb.close()
    with (OUT_DIR / "salesreturnlist.json").open("w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, default=str)
    print(f"Wrote {len(out)} rows -> {OUT_DIR / 'salesreturnlist.json'}")


if __name__ == "__main__":
    convert_sales()
    convert_returns()
