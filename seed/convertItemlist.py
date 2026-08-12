#!/usr/bin/env python3
"""
One-time convert Data/Itemlist.xlsx -> seed/data/itemlist.json

Requires openpyxl (pip3 install --user openpyxl).

Usage:
    python3 seed/convertItemlist.py [xlsx_path] [json_out_path]
    (defaults: Data/Itemlist.xlsx -> seed/data/itemlist.json)
"""
import json
import sys
from pathlib import Path

import openpyxl

DEFAULT_IN = "Data/Itemlist.xlsx"
DEFAULT_OUT = "seed/data/itemlist.json"


def main():
    in_path = Path(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_IN)
    out_path = Path(sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT)

    if not in_path.exists():
        print(f"Input file not found: {in_path}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(in_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]

    records = []
    for row in rows_iter:
        rec = dict(zip(header, row))
        records.append(rec)
    wb.close()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, default=str)

    print(f"Wrote {len(records)} rows -> {out_path}")


if __name__ == "__main__":
    main()
